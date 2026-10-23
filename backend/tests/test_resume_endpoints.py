import asyncio
import io
import docx
from pypdf import PdfWriter
import pytest

from httpx import AsyncClient, ASGITransport
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent.parent))

from main import app
from db.models import get_db, AsyncSessionLocal, Tenant, UserRole
from api.routes.auth import get_current_user, User
import uuid

# Mock auth user
mock_user = User(
    id=uuid.uuid4(),
    email="test.recruiter@dvt.ai",
    role=UserRole.ADMIN,
    tenant_id=uuid.uuid4(),
    full_name="Test Recruiter",
)


def create_sample_docx(test_email: str) -> bytes:
    doc = docx.Document()
    doc.add_heading("Jane Doe", level=0)
    doc.add_paragraph(f"Email: {test_email} | Phone: +1 415-555-0199 | Location: New York, NY")
    doc.add_paragraph("LinkedIn: https://linkedin.com/in/janedoe | GitHub: https://github.com/janedoe")
    doc.add_heading("Summary", level=1)
    doc.add_paragraph("Principal Distributed Systems Engineer with 8 years of production experience in Kubernetes, Go, and Kafka.")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Principal Systems Architect at CloudScale Tech (2021 - Present)")
    doc.add_paragraph("• Architected multi-region event streaming platform handling 2M events/sec.")
    doc.add_paragraph("• Reduced distributed query latency by 45% using Rust and RocksDB.")
    doc.add_paragraph("Senior Backend Engineer at Enterprise Corp (2018 - 2021)")
    doc.add_paragraph("• Designed and deployed 15 core microservices using Go and PostgreSQL.")
    doc.add_heading("Education", level=1)
    doc.add_paragraph("Bachelor of Science in Computer Science from MIT (2018)")
    doc.add_heading("Skills", level=1)
    doc.add_paragraph("Go, Python, Kubernetes, Kafka, Distributed Systems, Docker, PostgreSQL, AWS")
    
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_full_resume_flow():
    app.dependency_overrides[get_current_user] = lambda: mock_user
    try:
        # Ensure tenant exists in database
        async with AsyncSessionLocal() as session:
            t = Tenant(id=mock_user.tenant_id, name="Resume Flow Corp", domain=f"resume-{uuid.uuid4().hex[:6]}.com")
            session.add(t)
            await session.commit()

        test_email = f"jane.doe.{uuid.uuid4().hex[:8]}@example.com"
        docx_bytes = create_sample_docx(test_email)
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            # 1. Recruiter intake upload
            files = {
                "file": ("jane_doe_cv.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            }
            res = await client.post("/api/v1/candidates/upload", files=files)
            assert res.status_code == 200, res.text
            data = res.json()
            assert data["status"] == "success"
            cand = data["candidate"]
            cand_id = cand["id"]
            assert cand["first_name"] == "Jane"
            assert cand["last_name"] == "Doe"
            assert cand["email"] == test_email
            assert "Kubernetes" in cand["skills"]
            assert "Kafka" in cand["skills"]
            
            parsed = data["parsed_data"]
            assert len(parsed["experience"]) >= 1
            assert len(parsed["education"]) >= 1

            # 2. Get resume details
            details_res = await client.get(f"/api/v1/candidates/{cand_id}/resume/details")
            assert details_res.status_code == 200
            details = details_res.json()
            assert details["has_resume"] is True
            assert "Principal Distributed Systems Engineer" in details["raw_text"]
            assert "Architected multi-region event streaming" in details["raw_text"]

            # 3. Download original resume
            download_res = await client.get(f"/api/v1/candidates/{cand_id}/resume/download")
            assert download_res.status_code == 200
            assert download_res.content == docx_bytes
            assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in download_res.headers["content-type"]
    finally:
        app.dependency_overrides.pop(get_current_user, None)


if __name__ == "__main__":
    asyncio.run(test_full_resume_flow())
